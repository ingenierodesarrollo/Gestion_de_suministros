from decimal import Decimal
from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook

from areas.models import Area

from . import tests as base
from .consolidado import consolidado_por_area
from .exportar import NOMBRE_SIN_DATOS, _nombre_hoja
from .models import EstadoProyeccion, ProyeccionArea


class ConsolidadoCalculoTest(base.BaseProyeccionTest):

    def setUp(self):
        self.proy_lab = ProyeccionArea.objects.create(periodo=self.periodo, area=self.laboratorio, responsable=self.coord_lab)
        self.crear_linea(self.proy_urg, producto='Guantes', ene=10)
        self.crear_linea(self.proy_lab, producto='  GUANTES ', ene=5, feb=2)
        for p in (self.proy_urg, self.proy_lab):
            p.estado = EstadoProyeccion.ENVIADA
            p.save()

    def test_un_borrador_no_entra(self):
        self.proy_lab.estado = EstadoProyeccion.BORRADOR
        self.proy_lab.save()
        self.assertEqual(len(consolidado_por_area(self.periodo)), 1)

    def test_cada_area_va_por_separado(self):
        areas = consolidado_por_area(self.periodo)
        self.assertEqual([a['proyeccion'].area.nombre for a in areas], ['Laboratorio', 'Urgencias'])
        self.assertEqual(areas[1]['costo_total'], Decimal('15000'))
        self.assertEqual(len(areas[1]['filas']), 1)


class ConsolidadoPermisosYExcelTest(base.BaseProyeccionTest):

    def test_coordinador_de_area_no_entra(self):
        self.client.force_login(self.coord_urg)
        self.assertEqual(self.client.get(reverse('proyecciones:consolidado')).status_code, 403)
        self.assertEqual(self.client.get(reverse('proyecciones:descargar_consolidado', args=[2027])).status_code, 403)

    def test_administrador_ve_la_pantalla(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse('proyecciones:consolidado')).status_code, 200)

    def test_excel_trae_una_hoja_por_area(self):
        self.crear_linea(self.proy_urg, producto='Guantes')
        self.proy_urg.estado = EstadoProyeccion.ENVIADA
        self.proy_urg.save()
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('proyecciones:descargar_consolidado', args=[2027]))
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('spreadsheetml', respuesta['Content-Type'])
        libro = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(libro.sheetnames, ['Urgencias'])
        self.assertEqual(libro['Urgencias']['A7'].value, 'Guantes')

    def test_excel_sin_areas_enviadas_trae_hoja_sin_datos(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('proyecciones:descargar_consolidado', args=[2027]))
        libro = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(libro.sheetnames, [NOMBRE_SIN_DATOS])

    def test_anio_que_no_existe(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse('proyecciones:descargar_consolidado', args=[1999])).status_code, 404)

    def test_nombres_de_hoja_validos_y_sin_repetir(self):
        usadas = set()
        self.assertEqual(_nombre_hoja('Rayos X/Imágenes', usadas), 'Rayos XImágenes')
        self.assertEqual(_nombre_hoja('Rayos X/Imágenes', usadas), 'Rayos XImágenes_2')