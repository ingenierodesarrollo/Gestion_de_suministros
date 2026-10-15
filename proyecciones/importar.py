"""Carga de líneas desde Excel: la plantilla que se descarga y la lectura del archivo que sube el área.

La lectura busca las columnas por su NOMBRE, no por su posición. Así sirve tanto la plantilla
como un FR-GDS-11 de años anteriores (títulos en la fila 6, sin unidad ni precio)."""
import unicodedata
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .exportar import AZUL
from .models import AccionBitacora, Bitacora, Linea, UnidadMedida
from .servicios import registrar

MESES_NOMBRES = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio', 'Agosto',
                 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
# Mismo orden que el FR-GDS-11 (A producto, B dependencia, C serie, D-O meses); luego unidad y precio.
COLUMNAS = ['Equipo / Producto / Servicio', 'Dependencia', 'Serie', *MESES_NOMBRES, 'Unidad', 'Precio est.']
HOJA = 'Productos'
UNIDAD_POR_DEFECTO = 'Unidad'
PREFIJOS_IGNORADOS = ('EJEMPLO', 'TOTAL', 'SUBTOTAL', 'COSTO TOTAL')
FILAS_PARA_BUSCAR_TITULOS = 20
MAX_FILAS = 2000                 # filas con datos que se aceptan en un archivo
MAX_BYTES = 5 * 1024 * 1024      # 5 MB
MAX_CANTIDAD = 9_999_999         # tope por mes, para atrapar errores de digitación


def _clave(texto):
    """'Equipo / Producto / Servicio' -> 'equipoproductoservicio' (sin tildes, espacios ni signos)."""
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return ''.join(c for c in texto.lower() if c.isalnum())


# Nombres de columna que se aceptan para cada dato. Las demás columnas se ignoran
# (por ejemplo "Autorizado" y "Costo proyectado" del formato original).
ALIAS = {
    'producto': {'equipoproductoservicio', 'productoequiposervicio', 'producto', 'productos',
                 'equipo', 'servicio', 'productoservicio', 'descripcion'},
    'dependencia': {'dependencia'},
    'serie': {'serie'},
    'unidad': {'unidad', 'unidaddemedida', 'presentacion'},
    'precio': {'precioest', 'precio', 'preciounitario', 'precioestimado', 'preciounitarioestimado',
               'preciounitarioest'},
}
for _campo, _nombre in zip(Linea.MESES, MESES_NOMBRES):
    ALIAS[_campo] = {_clave(_nombre), _campo}
ALIAS['sep'].update({'sept', 'setiembre'})
CAMPO_DE_ALIAS = {alias: campo for campo, nombres in ALIAS.items() for alias in nombres}


class ErrorArchivo(Exception):
    """El archivo completo no sirve (no es Excel, sin columna de producto, demasiadas filas)."""


# ---------- Plantilla ----------

def construir_plantilla(unidades):
    """Devuelve los bytes de la plantilla: hoja Productos con un ejemplo y hoja Unidades con las válidas."""
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    for col, titulo in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col, value=titulo)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=AZUL)
        celda.alignment = Alignment(horizontal='center', wrap_text=True)
    # Fila de ejemplo: al leer, toda fila cuyo producto empiece con EJEMPLO se ignora.
    ws.append(['EJEMPLO - Guantes de nitrilo talla M', '', '', *[10] * 12, UNIDAD_POR_DEFECTO, 1500])
    for letra, ancho in {'A': 45, 'B': 18, 'C': 14, 'P': 18, 'Q': 12}.items():
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = 'B2'

    hoja_unidades = wb.create_sheet('Unidades')
    hoja_unidades['A1'] = 'Unidades válidas (si deja la unidad vacía, queda "Unidad")'
    hoja_unidades['A1'].font = Font(bold=True, color=AZUL)
    for fila, nombre in enumerate(unidades, start=2):
        hoja_unidades.cell(row=fila, column=1, value=nombre)
    hoja_unidades.column_dimensions['A'].width = 55

    if unidades:
        # Lista desplegable en la columna Unidad (P), para no escribirla a mano.
        lista = DataValidation(type='list', formula1=f'=Unidades!$A$2:$A${len(unidades) + 1}', allow_blank=True)
        lista.error = 'Elija una unidad de la lista (hoja Unidades) o déjela vacía.'
        lista.errorTitle = 'Unidad no válida'
        ws.add_data_validation(lista)
        lista.add(f'P2:P{MAX_FILAS + 1}')

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ---------- Lectura ----------

def _texto(valor):
    return '' if valor is None else str(valor).strip()


def _entero(valor):
    """Cantidad de un mes. Vacío = None (en una línea nueva cuenta como 0; en una existente no cambia nada).
    Debe ser entero y no negativo."""
    if _texto(valor) == '':
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    if isinstance(valor, int) and not isinstance(valor, bool):
        numero = valor
    elif _texto(valor).isdigit():
        numero = int(_texto(valor))
    else:
        raise ValueError
    if numero < 0 or numero > MAX_CANTIDAD:
        raise ValueError
    return numero


def _precio(valor):
    """Precio unitario estimado: opcional (vacío = None); si viene, mayor que cero. Acepta 1500, 1500.5 o 1.500,50."""
    texto = _texto(valor).replace('$', '').replace(' ', '')
    if not texto:
        return None  # se completa después; no se deja enviar la proyección sin precio
    if isinstance(valor, (int, float)):
        texto = str(valor)
    elif ',' in texto:
        texto = texto.replace('.', '').replace(',', '.')
    try:
        precio = Decimal(texto).quantize(Decimal('0.01'))
    except InvalidOperation:
        raise ValueError(f'el precio "{_texto(valor)}" no es un número')
    if precio <= 0 or precio >= Decimal('10000000000'):
        raise ValueError('el precio debe ser mayor que cero')
    return precio


def unidad_por_defecto():
    """La unidad "Unidad", que se usa cuando una línea llega sin unidad. Si no existe, la crea."""
    unidad = UnidadMedida.objects.filter(nombre__iexact=UNIDAD_POR_DEFECTO).first()
    return unidad or UnidadMedida.objects.create(nombre=UNIDAD_POR_DEFECTO)


def unidades_validas():
    """{'caja x 100': (id, 'Caja x 100'), ...} para buscar la unidad sin importar mayúsculas."""
    unidad_por_defecto()  # asegura que exista para las filas sin unidad
    return {u.nombre.strip().lower(): (u.pk, u.nombre) for u in UnidadMedida.objects.all()}


def _clave_linea(producto, dependencia, serie, unidad_id):
    """Dos filas son la misma línea solo si coinciden producto, dependencia, serie y unidad.
    Así "Guantes M" de dos dependencias, o dos monitores con series distintas, no se toman por repetidos."""
    return (' '.join(producto.lower().split()), dependencia.strip().lower(), serie.strip().lower(), unidad_id)


def lineas_existentes(proyeccion):
    """{clave: datos actuales} de las líneas que ya tiene la proyección. Sirve para saber cuáles se actualizan."""
    campos = ('id', 'producto', 'dependencia', 'serie', 'unidad_id', *Linea.MESES, 'precio_unitario_est')
    existentes = {}
    for fila in proyeccion.lineas.values_list(*campos):
        datos = dict(zip(campos, fila))
        clave = _clave_linea(datos['producto'], datos['dependencia'], datos['serie'], datos['unidad_id'])
        existentes[clave] = datos
    return existentes


def _cambios(actual, meses, precio):
    """Lista de (campo, antes, después) solo con las celdas que el Excel trae con valor y que son distintas."""
    cambios = []
    for campo, valor in zip(Linea.MESES, meses):
        if valor is not None and valor != actual[campo]:
            cambios.append((campo, actual[campo], valor))
    if precio is not None:
        antes = actual['precio_unitario_est']
        if antes is None or precio != antes:
            cambios.append(('precio_unitario_est', None if antes is None else str(antes), str(precio)))
    return cambios


def _columnas(valores):
    """Si la fila es la de títulos, devuelve {campo: posición}; si no, None."""
    posiciones = {}
    for i, valor in enumerate(valores):
        campo = CAMPO_DE_ALIAS.get(_clave(valor))
        if campo and campo not in posiciones:
            posiciones[campo] = i
    return posiciones if 'producto' in posiciones else None


def _buscar_titulos(wb):
    """Busca la hoja y la fila de títulos. Prefiere la hoja Productos de la plantilla."""
    hojas = sorted(wb.worksheets, key=lambda h: h.title != HOJA)
    for ws in hojas:
        for numero, valores in enumerate(ws.iter_rows(max_row=FILAS_PARA_BUSCAR_TITULOS, values_only=True), start=1):
            posiciones = _columnas(valores)
            if posiciones:
                return ws, numero, posiciones
    raise ErrorArchivo('No se encontró la columna "Equipo / Producto / Servicio". '
                       'Use la plantilla o el formato FR-GDS-11.')


def leer_excel(archivo, unidades, existentes):
    """Lee el archivo y separa filas buenas, filas con errores y filas repetidas. No guarda nada."""
    if getattr(archivo, 'size', 0) > MAX_BYTES:
        raise ErrorArchivo('El archivo pesa más de 5 MB.')
    try:
        wb = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        raise ErrorArchivo('No se pudo abrir el archivo. Debe ser Excel guardado como .xlsx.')
    try:
        return _leer_hoja(*_buscar_titulos(wb), unidades, existentes)
    finally:
        wb.close()


def _leer_hoja(ws, fila_titulos, posiciones, unidades, existentes):
    resultado = {'filas': [], 'actualizar': [], 'sin_cambios': 0, 'errores': [], 'repetidas': []}
    vistos = set()
    con_datos = 0
    for numero, valores in enumerate(ws.iter_rows(min_row=fila_titulos + 1, values_only=True), start=fila_titulos + 1):
        def dato(campo):
            i = posiciones.get(campo)
            return valores[i] if i is not None and i < len(valores) else None

        producto, dependencia, serie, unidad = (_texto(dato(c)) for c in ('producto', 'dependencia', 'serie', 'unidad'))
        if not any(_texto(dato(c)) for c in posiciones):
            continue  # fila vacía
        if producto.upper().startswith(PREFIJOS_IGNORADOS):
            continue  # ejemplo de la plantilla o fila de totales del Excel viejo
        con_datos += 1
        if con_datos > MAX_FILAS:
            raise ErrorArchivo(f'El archivo tiene más de {MAX_FILAS} productos. Divídalo en dos archivos.')

        problemas = []
        if not producto:
            problemas.append('falta el equipo / producto / servicio')
        elif len(producto) > 150:
            problemas.append('el producto tiene más de 150 caracteres')
        if len(dependencia) > 100:
            problemas.append('la dependencia tiene más de 100 caracteres')
        if len(serie) > 100:
            problemas.append('la serie tiene más de 100 caracteres')
        encontrada = unidades.get((unidad or UNIDAD_POR_DEFECTO).lower())
        if encontrada is None:
            problemas.append(f'la unidad "{unidad}" no existe')
        meses = []
        for campo, nombre in zip(Linea.MESES, MESES_NOMBRES):
            try:
                meses.append(_entero(dato(campo)))
            except ValueError:
                problemas.append(f'{nombre}: "{_texto(dato(campo))}" no es una cantidad entera válida')
                meses.append(None)
        precio = None
        try:
            precio = _precio(dato('precio'))
        except ValueError as error:
            problemas.append(str(error))

        if problemas:
            resultado['errores'].append((numero, producto or '—', '; '.join(problemas)))
            continue
        clave = _clave_linea(producto, dependencia, serie, encontrada[0])
        if clave in vistos:  # la misma línea dos veces en el archivo
            resultado['repetidas'].append((numero, producto, encontrada[1], dependencia, serie))
            continue
        vistos.add(clave)
        actual = existentes.get(clave)
        if actual is None:  # línea nueva: los meses vacíos cuentan como 0
            meses_nuevos = [m or 0 for m in meses]
            resultado['filas'].append({
                'fila': numero, 'producto': producto, 'dependencia': dependencia, 'serie': serie,
                'unidad_id': encontrada[0], 'unidad': encontrada[1], 'meses': meses_nuevos,
                'total': sum(meses_nuevos), 'precio': None if precio is None else str(precio),  # texto, para la sesión
            })
            continue
        cambios = _cambios(actual, meses, precio)  # línea que ya existe: solo cambian las celdas con valor
        if cambios:
            resultado['actualizar'].append({
                'fila': numero, 'id': actual['id'], 'producto': producto, 'dependencia': dependencia,
                'serie': serie, 'unidad': encontrada[1], 'cambios': cambios,
            })
        else:
            resultado['sin_cambios'] += 1
    return resultado


# ---------- Guardar ----------

def nombre_campo(campo):
    """'ene' -> 'Enero'; 'precio_unitario_est' -> 'Precio est.'."""
    return dict(zip(Linea.MESES, MESES_NOMBRES)).get(campo, 'Precio est.')


def describir_cambios(cambios):
    """'Enero: 5 → 8; Precio est.: 1500.00 → 1800.50' para mostrar en la vista previa."""
    return '; '.join(f'{nombre_campo(campo)}: {"vacío" if antes is None else antes} → {despues}'
                     for campo, antes, despues in cambios)


@transaction.atomic
def aplicar_importacion(proyeccion, nuevas, actualizar, usuario):
    """Crea las líneas nuevas y actualiza las existentes que el usuario confirmó en la vista previa.
    Deja rastro en la bitácora. Devuelve (creadas, actualizadas)."""
    existentes = lineas_existentes(proyeccion)
    crear = []
    for f in nuevas:
        clave = _clave_linea(f['producto'], f['dependencia'], f['serie'], f['unidad_id'])
        if clave in existentes:
            continue  # alguien la creó mientras tanto: no se duplica
        existentes[clave] = None
        crear.append(Linea(
            proyeccion=proyeccion, producto=f['producto'], dependencia=f['dependencia'], serie=f['serie'],
            unidad_id=f['unidad_id'], precio_unitario_est=Decimal(f['precio']) if f['precio'] else None,
            **dict(zip(Linea.MESES, f['meses'])),
        ))
    Linea.objects.bulk_create(crear, batch_size=500)

    # Solo se tocan líneas de ESTA proyección, aunque la sesión traiga otros números.
    lineas = Linea.objects.in_bulk([a['id'] for a in actualizar]) if actualizar else {}
    rastro = []
    modificadas = []
    for a in actualizar:
        linea = lineas.get(a['id'])
        if linea is None or linea.proyeccion_id != proyeccion.pk:
            continue
        for campo, antes, despues in a['cambios']:
            setattr(linea, campo, Decimal(despues) if campo == 'precio_unitario_est' else despues)
            rastro.append(Bitacora(
                usuario=usuario, accion=AccionBitacora.EDITAR, tabla='linea', registro_id=linea.pk,
                campo=campo, valor_anterior='' if antes is None else str(antes)[:200],
                valor_nuevo=str(despues)[:200], motivo='Carga desde Excel',
            ))
        modificadas.append(linea)
    if modificadas:
        Linea.objects.bulk_update(modificadas, [*Linea.MESES, 'precio_unitario_est'], batch_size=500)
        Bitacora.objects.bulk_create(rastro, batch_size=500)
        registrar(usuario, AccionBitacora.EDITAR, proyeccion, 'líneas actualizadas desde Excel', None, len(modificadas))
    if crear:
        registrar(usuario, AccionBitacora.CREAR, proyeccion, 'líneas desde Excel', None, len(crear))
    return len(crear), len(modificadas)
