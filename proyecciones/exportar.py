"""Genera el Excel del consolidado: una hoja por área y un resumen por producto."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .consolidado import MESES_CORTOS, consolidado_por_area, consolidado_por_producto

AZUL = '172554'
AZUL_CLARO = 'DBEAFE'
BORDE = Border(*(Side(style='thin', color='C3C8D4'),) * 4)
NOMBRE_RESUMEN = 'Resumen por producto'
FORMATO_DINERO = '#,##0.00'


def _nombre_hoja(nombre, usadas):
    """Excel limita el nombre de hoja a 31 letras y no admite ciertos signos."""
    base = ''.join(c for c in nombre if c not in '[]:*?/\\')[:26] or 'Area'
    nombre_hoja, n = base, 2
    while nombre_hoja.lower() in usadas:
        nombre_hoja = f'{base}_{n}'
        n += 1
    usadas.add(nombre_hoja.lower())
    return nombre_hoja


def _encabezado(ws, periodo, linea_4, columnas):
    """Las primeras filas, parecidas al formato FR-GDS-11."""
    textos = ['GESTIÓN DE SUMINISTROS', 'PROYECCIÓN DE COMPRAS',
              f'Código: FR-GDS-11 · AÑO: {periodo.anio}', linea_4]
    for i, texto in enumerate(textos, start=1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=columnas)
        celda = ws.cell(row=i, column=1, value=texto)
        celda.font = Font(bold=True, size=14 if i == 1 else 11, color=AZUL)
        celda.alignment = Alignment(horizontal='center')


def _titulos_columnas(ws, fila, titulos):
    for col, titulo in enumerate(titulos, start=1):
        c = ws.cell(row=fila, column=col, value=titulo)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AZUL)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDE


def _escribir_fila(ws, fila, valores, num_desde, dinero_en):
    """Escribe una fila; los valores desde la columna num_desde van alineados a la derecha."""
    for col, valor in enumerate(valores, start=1):
        c = ws.cell(row=fila, column=col, value=valor)
        c.border = BORDE
        if col >= num_desde:
            c.alignment = Alignment(horizontal='right')
        if col in dinero_en:
            c.number_format = FORMATO_DINERO


def _anchos(ws, fijos, columnas):
    for col in range(1, columnas + 1):
        ws.column_dimensions[get_column_letter(col)].width = fijos.get(col, 7)


def _hoja_area(wb, periodo, datos, usadas):
    p = datos['proyeccion']
    ws = wb.create_sheet(_nombre_hoja(p.area.nombre, usadas))
    columnas = 20
    _encabezado(ws, periodo, f'Área: {p.area.nombre} · Estado: {p.get_estado_display()}', columnas)
    _titulos_columnas(ws, 6, ['Equipo / Producto / Servicio', 'Dependencia', 'Serie', 'Unidad']
                      + MESES_CORTOS + ['Total año', 'Autorizado', 'Precio est.', 'Costo proyectado'])
    fila = 7
    for f in datos['filas']:
        l = f['linea']
        _escribir_fila(ws, fila, [l.producto, l.dependencia, l.serie, l.unidad.nombre] + f['meses']
                       + [l.total_anual, l.cantidad_autorizada, l.precio_unitario_est, l.costo],
                       num_desde=5, dinero_en={19, 20})
        fila += 1
    ws.cell(row=fila, column=1, value='Costo total del área').font = Font(bold=True)
    total = ws.cell(row=fila, column=20, value=datos['costo_total'])
    total.font = Font(bold=True, color=AZUL)
    total.number_format = FORMATO_DINERO
    total.fill = PatternFill('solid', fgColor=AZUL_CLARO)
    _anchos(ws, {1: 32, 2: 18, 3: 14, 4: 14, 17: 10, 18: 11, 19: 12, 20: 16}, columnas)
    ws.freeze_panes = 'E7'


def _hoja_resumen(wb, periodo, productos):
    ws = wb.create_sheet(NOMBRE_RESUMEN)
    columnas = 18
    _encabezado(ws, periodo, 'Resumen por producto (suma de todas las áreas enviadas y aprobadas)', columnas)
    _titulos_columnas(ws, 6, ['Equipo / Producto / Servicio', 'Unidad', 'Áreas']
                      + MESES_CORTOS + ['Total año', 'Autorizado', 'Costo proyectado'])
    fila = 7
    for g in productos:
        _escribir_fila(ws, fila, [g['producto'], g['unidad'], g['areas']] + g['meses']
                       + [g['total'], g['autorizado'], g['costo']], num_desde=4, dinero_en={18})
        fila += 1
    ws.cell(row=fila, column=1, value='Costo total').font = Font(bold=True)
    total = ws.cell(row=fila, column=18, value=sum(g['costo'] for g in productos))
    total.font = Font(bold=True, color=AZUL)
    total.number_format = FORMATO_DINERO
    total.fill = PatternFill('solid', fgColor=AZUL_CLARO)
    _anchos(ws, {1: 32, 2: 14, 3: 28, 16: 10, 17: 11, 18: 16}, columnas)
    ws.freeze_panes = 'D7'


def construir_excel(periodo):
    """Devuelve los bytes del archivo .xlsx del consolidado del periodo."""
    wb = Workbook()
    wb.remove(wb.active)
    usadas = {NOMBRE_RESUMEN.lower()}
    for datos in consolidado_por_area(periodo):
        _hoja_area(wb, periodo, datos, usadas)
    _hoja_resumen(wb, periodo, consolidado_por_producto(periodo))
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()