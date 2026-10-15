"""Pruebas de la carga de líneas desde Excel. Solo datos inventados."""
from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .importar import construir_plantilla
from .models import Bitacora, EstadoProyeccion, Linea
from .tests import BaseProyeccionTest


def excel(*filas):
    """Arma un .xlsx con la plantilla y las filas dadas, listo para subir."""
    wb = load_workbook(BytesIO(construir_plantilla(['Unidad'])))
    hoja = wb['Productos']
    for fila in filas:
        hoja.append(list(fila))
    salida = BytesIO()
    wb.save(salida)
    return SimpleUploadedFile('carga.xlsx', salida.getvalue())


MESES_VACIOS = [None] * 12
# Orden de la plantilla: producto, dependencia, serie, 12 meses, unidad, precio.


class PlantillaTest(BaseProyeccionTest):

    def test_coordinador_descarga_plantilla(self):
        self.client.force_login(self.coord_urg)
        respuesta = self.client.get(reverse('proyecciones:plantilla_excel'))
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('attachment', respuesta['Content-Disposition'])


class ImportarExcelTest(BaseProyeccionTest):

    def setUp(self):
        self.client.force_login(self.coord_urg)
        self.url = reverse('proyecciones:importar_excel', args=[self.proy_urg.pk])
        self.url_confirmar = reverse('proyecciones:confirmar_importacion', args=[self.proy_urg.pk])

    def subir(self, *filas):
        return self.client.post(self.url, {'archivo': excel(*filas)})

    def test_vista_previa_no_guarda(self):
        respuesta = self.subir(['Jeringa 5 ml', '', '', 1, 2, *[0] * 10, 'Unidad', 500])
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, 'Confirmar y agregar 1 líneas')
        self.assertEqual(self.proy_urg.lineas.count(), 0)

    def test_confirmar_crea_lineas_y_bitacora(self):
        self.subir(['Jeringa 5 ml', '', '', 1, 2, *[0] * 10, 'Unidad', '1.250,50'],
                   ['Gasa', 'Sala 1', '', *MESES_VACIOS, 'unidad', 300])
        self.client.post(self.url_confirmar)
        self.assertEqual(self.proy_urg.lineas.count(), 2)
        jeringa = self.proy_urg.lineas.get(producto='Jeringa 5 ml')
        self.assertEqual(jeringa.total_anual, 3)
        self.assertEqual(jeringa.precio_unitario_est, Decimal('1250.50'))
        self.assertTrue(Bitacora.objects.filter(campo='líneas desde Excel', valor_nuevo='2').exists())

    def test_solo_nombres_sin_precio_ni_cantidades(self):
        self.subir(['Monitor de signos vitales', '', '', *MESES_VACIOS, None, None])
        self.client.post(self.url_confirmar)
        linea = self.proy_urg.lineas.get()
        self.assertIsNone(linea.precio_unitario_est)
        self.assertEqual(linea.total_anual, 0)
        self.assertEqual(linea.unidad, self.unidad)  # sin unidad queda en "Unidad"

    def test_errores_y_repetidas_no_se_cargan(self):
        self.crear_linea(self.proy_urg, producto='Guantes')
        self.subir(['GUANTES', '', '', *MESES_VACIOS, 'Unidad', None],      # ya existe
                   ['Bata', '', '', *MESES_VACIOS, 'Paquete', None],        # unidad que no existe
                   ['Lancetas', '', '', -3, *[0] * 11, 'Unidad', None],     # cantidad negativa
                   ['Alcohol', '', '', *MESES_VACIOS, 'Unidad', 'barato'])  # precio inválido
        self.client.post(self.url_confirmar)
        self.assertEqual(self.proy_urg.lineas.count(), 1)  # solo la que ya estaba

    def test_archivo_que_no_es_excel(self):
        archivo = SimpleUploadedFile('carga.xlsx', b'esto no es un excel')
        respuesta = self.client.post(self.url, {'archivo': archivo})
        self.assertContains(respuesta, 'No se pudo abrir el archivo')

    def test_confirmar_dos_veces_no_duplica(self):
        self.subir(['Gasa', '', '', *MESES_VACIOS, 'Unidad', 300])
        self.client.post(self.url_confirmar)
        self.client.post(self.url_confirmar)
        self.assertEqual(self.proy_urg.lineas.count(), 1)

    def test_formato_original_fr_gds_11(self):
        """Un Excel de años anteriores: títulos en la fila 6, sin unidad ni precio, con Autorizado y Costo."""
        wb = Workbook()
        hoja = wb.active
        hoja.title = 'servicios'
        hoja['A1'] = 'GESTION DE SUMINISTROS'
        titulos = ['Equipo/Producto/Servicio', 'Dependencia', 'Serie', 'Enero', 'Febrero', 'Marzo', 'Abril',
                   'Mayo', 'Junio', 'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre',
                   'Autorizado', 'Costo proyectado']
        for col, titulo in enumerate(titulos, start=1):
            hoja.cell(row=6, column=col, value=titulo)
        hoja.cell(row=7, column=1, value='Mantenimiento aire acondicionado')
        hoja.cell(row=7, column=4, value=1)
        hoja.cell(row=7, column=16, value='1/1')
        hoja.cell(row=7, column=17, value=500000)
        salida = BytesIO()
        wb.save(salida)
        self.client.post(self.url, {'archivo': SimpleUploadedFile('2025.xlsx', salida.getvalue())})
        self.client.post(self.url_confirmar)
        linea = self.proy_urg.lineas.get()
        self.assertEqual(linea.producto, 'Mantenimiento aire acondicionado')
        self.assertEqual(linea.ene, 1)
        self.assertIsNone(linea.precio_unitario_est)
        self.assertIsNone(linea.cantidad_autorizada)  # el autorizado lo ponen los administradores

    def test_otra_area_no_puede(self):
        self.client.force_login(self.coord_lab)
        self.assertEqual(self.client.get(self.url).status_code, 403)

    def test_enviada_no_se_carga(self):
        self.proy_urg.estado = EstadoProyeccion.ENVIADA
        self.proy_urg.save()
        self.assertEqual(self.client.get(self.url).status_code, 403)


class EnviarIncompletaTest(BaseProyeccionTest):

    def setUp(self):
        self.client.force_login(self.coord_urg)
        self.url = reverse('proyecciones:enviar', args=[self.proy_urg.pk])

    def test_no_envia_con_linea_sin_precio(self):
        self.crear_linea(self.proy_urg, precio_unitario_est=None)
        self.client.post(self.url)
        self.proy_urg.refresh_from_db()
        self.assertEqual(self.proy_urg.estado, EstadoProyeccion.BORRADOR)

    def test_no_envia_con_linea_sin_cantidades(self):
        self.crear_linea(self.proy_urg, ene=0)
        self.client.post(self.url)
        self.proy_urg.refresh_from_db()
        self.assertEqual(self.proy_urg.estado, EstadoProyeccion.BORRADOR)


class LineaSinUnidadTest(BaseProyeccionTest):

    def test_agregar_linea_sin_unidad_queda_en_unidad(self):
        self.client.force_login(self.coord_urg)
        datos = {'producto': 'Guantes', 'dependencia': '', 'serie': '', 'unidad': '',
                 'precio_unitario_est': '', **{mes: 0 for mes in Linea.MESES}}
        self.client.post(reverse('proyecciones:agregar_linea', args=[self.proy_urg.pk]), datos)
        linea = self.proy_urg.lineas.get()
        self.assertEqual(linea.unidad, self.unidad)
        self.assertIsNone(linea.precio_unitario_est)
